"""
Document chunking pipeline — Section 2, Step 1.

Supports: PDF, Markdown, plain text.
Produces: list[DocumentChunk] with rich metadata.
"""
from __future__ import annotations

import hashlib
import re
import uuid
from pathlib import Path
from typing import Any, Optional

from langchain_text_splitters import RecursiveCharacterTextSplitter
from pydantic import BaseModel

from src.config import get_settings
from src.models.schemas import DocumentChunk
from src.utils.logging import get_logger

log = get_logger(__name__)


# ──────────────────────────────────────────────────────────────────────────────
# Loader helpers
# ──────────────────────────────────────────────────────────────────────────────

def _load_pdf(path: Path) -> list[tuple[str, dict[str, Any]]]:
    """Load a PDF and return (page_text, metadata) pairs per page."""
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        pages.append((text, {"page_number": i + 1, "total_pages": len(reader.pages)}))
    return pages


def _load_markdown(path: Path) -> list[tuple[str, dict[str, Any]]]:
    """Load a Markdown file, preserving heading context as metadata."""
    content = path.read_text(encoding="utf-8")
    # Extract h1/h2 headings for section metadata
    sections: list[tuple[str, dict[str, Any]]] = []
    current_heading = ""
    buffer: list[str] = []

    for line in content.splitlines(keepends=True):
        if re.match(r"^#{1,2} ", line):
            if buffer:
                sections.append(("".join(buffer), {"section": current_heading}))
                buffer = []
            current_heading = line.strip("# \n")
        buffer.append(line)

    if buffer:
        sections.append(("".join(buffer), {"section": current_heading}))

    return sections if sections else [(content, {})]


def _load_docx(path: Path) -> list[tuple[str, dict[str, Any]]]:
    """Load a DOCX file, preserving paragraph headings as metadata."""
    try:
        from docx import Document
        doc = Document(str(path))
        sections: list[tuple[str, dict[str, Any]]] = []
        current_heading = ""
        buffer: list[str] = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            if p.style and p.style.name and p.style.name.startswith("Heading"):
                if buffer:
                    sections.append(("\n".join(buffer), {"section": current_heading}))
                    buffer = []
                current_heading = text
            buffer.append(text)
        if buffer:
            sections.append(("\n".join(buffer), {"section": current_heading}))
        return sections if sections else [("\n".join(p.text for p in doc.paragraphs if p.text), {})]
    except Exception as exc:
        log.warning("docx_load_failed_falling_back_to_text", error=str(exc))
        return _load_text(path)


def _load_text(path: Path) -> list[tuple[str, dict[str, Any]]]:
    return [(path.read_text(encoding="utf-8"), {})]


# ──────────────────────────────────────────────────────────────────────────────
# Main chunker
# ──────────────────────────────────────────────────────────────────────────────

class DocumentChunker:
    """
    Splits enterprise documents into semantically coherent chunks.

    Deterministic chunk IDs based on (source_uri, chunk_index) make
    re-ingestion idempotent — chunks can be upserted safely.
    """

    def __init__(
        self,
        chunk_size: Optional[int] = None,
        chunk_overlap: Optional[int] = None,
    ) -> None:
        settings = get_settings()
        self._splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size or settings.chunk_size,
            chunk_overlap=chunk_overlap or settings.chunk_overlap,
            length_function=len,
            separators=["\n\n", "\n", ". ", "! ", "? ", " ", ""],
        )

    def chunk_document(
        self,
        source_uri: str,
        document_title: Optional[str] = None,
        extra_metadata: Optional[dict[str, Any]] = None,
    ) -> list[DocumentChunk]:
        """
        Entry point. Detects file type, loads, splits, and packages chunks.
        """
        path = Path(source_uri)
        if not path.exists():
            raise FileNotFoundError(f"Document not found: {source_uri}")

        suffix = path.suffix.lower()
        loaders = {".pdf": _load_pdf, ".docx": _load_docx, ".md": _load_markdown, ".txt": _load_text}
        loader = loaders.get(suffix, _load_text)

        raw_sections = loader(path)
        log.info(
            "document_loaded",
            source_uri=source_uri,
            sections=len(raw_sections),
            file_type=suffix,
        )

        chunks: list[DocumentChunk] = []
        global_index = 0

        for section_text, section_meta in raw_sections:
            if not section_text.strip():
                continue

            splits = self._splitter.split_text(section_text)

            for split_text in splits:
                if not split_text.strip():
                    continue

                # Deterministic UUID from (uri, index) — idempotent re-ingestion
                chunk_id = uuid.UUID(
                    hashlib.md5(
                        f"{source_uri}::{global_index}".encode()
                    ).hexdigest()
                )

                metadata: dict[str, Any] = {
                    "file_type": suffix,
                    **(section_meta or {}),
                    **(extra_metadata or {}),
                }

                chunks.append(
                    DocumentChunk(
                        chunk_id=chunk_id,
                        source_uri=source_uri,
                        document_title=document_title or path.stem,
                        chunk_index=global_index,
                        chunk_text=split_text,
                        token_count=len(split_text.split()),  # rough token proxy
                        metadata=metadata,
                    )
                )
                global_index += 1

        log.info(
            "document_chunked",
            source_uri=source_uri,
            total_chunks=len(chunks),
        )
        return chunks
